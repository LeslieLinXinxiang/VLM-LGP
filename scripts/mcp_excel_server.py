#!/usr/bin/env python3
"""Local Excel MCP server (stdio transport).

Tools:
- list_sheets(file_path)
- read_sheet(file_path, sheet_name, max_rows=200, max_cols=50)
- write_cell(file_path, sheet_name, cell, value)
- append_row(file_path, sheet_name, values)

Supports .xlsx / .xlsm via openpyxl.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

try:
    from fastmcp import FastMCP  # type: ignore[import-not-found]
except Exception:
    from mcp.server.fastmcp import FastMCP  # type: ignore[import-not-found]

from openpyxl import load_workbook  # type: ignore[import-not-found]

mcp = FastMCP("excel-local")


ALLOWED_SUFFIXES = {".xlsx", ".xlsm"}


def _normalize_path(file_path: str) -> Path:
    p = Path(file_path).expanduser().resolve()
    if not p.exists():
        raise FileNotFoundError(f"Excel file not found: {p}")
    if p.suffix.lower() not in ALLOWED_SUFFIXES:
        raise ValueError("Only .xlsx / .xlsm files are supported")
    return p


@mcp.tool()
def list_sheets(file_path: str) -> list[str]:
    """List sheet names in an Excel file."""
    p = _normalize_path(file_path)
    wb = load_workbook(filename=p, data_only=False)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


@mcp.tool()
def read_sheet(
    file_path: str,
    sheet_name: str,
    max_rows: int = 200,
    max_cols: int = 50,
) -> dict[str, Any]:
    """Read a sheet as a 2D list (bounded by max_rows/max_cols)."""
    p = _normalize_path(file_path)
    wb = load_workbook(filename=p, data_only=True)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")
        ws = wb[sheet_name]

        max_rows = max(1, int(max_rows))
        max_cols = max(1, int(max_cols))

        rows: list[list[Any]] = []
        for r in ws.iter_rows(min_row=1, max_row=max_rows, min_col=1, max_col=max_cols, values_only=True):
            rows.append(list(r))

        return {
            "file": str(p),
            "sheet": sheet_name,
            "rows": rows,
            "max_rows": max_rows,
            "max_cols": max_cols,
        }
    finally:
        wb.close()


@mcp.tool()
def write_cell(file_path: str, sheet_name: str, cell: str, value: Any) -> dict[str, Any]:
    """Write a single cell value, e.g. cell='B2'."""
    p = _normalize_path(file_path)
    wb = load_workbook(filename=p)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")
        ws = wb[sheet_name]
        ws[cell] = value
        wb.save(p)
        return {
            "ok": True,
            "file": str(p),
            "sheet": sheet_name,
            "cell": cell,
            "value": value,
        }
    finally:
        wb.close()


@mcp.tool()
def append_row(file_path: str, sheet_name: str, values: list[Any]) -> dict[str, Any]:
    """Append one row to the bottom of a sheet."""
    p = _normalize_path(file_path)
    wb = load_workbook(filename=p)
    try:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"Sheet not found: {sheet_name}")
        ws = wb[sheet_name]
        ws.append(values)
        wb.save(p)
        return {
            "ok": True,
            "file": str(p),
            "sheet": sheet_name,
            "appended_values": values,
            "new_row_index": ws.max_row,
        }
    finally:
        wb.close()


if __name__ == "__main__":
    mcp.run()
